"""上岸鸭新增支持：PPT、Excel；自愿启用付费 API 视频语音转写。"""
import os,io,re,json,subprocess,tempfile,requests,threading
from pathlib import Path
from urllib.parse import urlparse
from rescue_service import RescueError, config_snapshot
_TRANS_LOCK=threading.RLock()
_TRANSCRIBE={'key':'','base':'https://api.openai.com/v1','model':'gpt-4o-mini-transcribe','enabled':False}

def parse_slide_deck(raw):
    from pptx import Presentation
    prs=Presentation(io.BytesIO(raw))
    segments,outline=[],[]
    if len(prs.slides)>300:raise RescueError('PPT超过300页，请拆分上传','PPT_TOO_LONG',413)
    for i,slide in enumerate(prs.slides,1):
        lines=[]
        for shape in slide.shapes:
            if getattr(shape,'has_text_frame',False):
                v=shape.text.strip()
                if v:lines.append(v)
            if getattr(shape,'has_table',False):
                for row in shape.table.rows:
                    lines.append(' | '.join(str(c.text).strip() for c in row.cells))
        notes=getattr(slide,'notes_slide',None)
        if notes:
            try:
                note=notes.notes_text_frame.text.strip()
                if note:lines.append('【演讲备注】'+note[:3000])
            except Exception:pass
        text='\n'.join(lines).strip()
        title=lines[0].split('\n',1)[0][:150] if lines else '第'+str(i)+'页'
        outline.append({'level':1,'title':'第'+str(i)+'页：'+title,'content':text[:10000]})
        segments.append('# 第'+str(i)+'页：'+title+'\n'+text)
    return '\n\n'.join(segments),outline

def parse_workbook(raw):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    sections,outline=[],[]
    for sh in wb.worksheets[:24]:
        lines=[]
        for ridx,row in enumerate(sh.iter_rows(max_row=min(sh.max_row or 0,1500),max_col=min(sh.max_column or 0,60),values_only=True)):
            contents=[str(v).strip().replace('\n',' ')[:300] if v is not None else '' for v in row]
            if any(contents):lines.append(' | '.join(contents))
            if len(lines)>=500:break
        text='\n'.join(lines)
        outline.append({'level':1,'title':'工作表：'+sh.title,'content':text[:15000]})
        sections.append('# 工作表：'+sh.title+'\n'+text)
    wb.close()
    return '\n\n'.join(sections),outline

def parse_extra(name,raw):
    ext=Path(name).suffix.lower()
    try:
        if ext in ('.pptx','.pptm'):text,outline=parse_slide_deck(raw)
        elif ext in ('.xlsx','.xlsm'):text,outline=parse_workbook(raw)
        else:raise RescueError('不支持的表格或演示文件，请另存为 .pptx 或 .xlsx。','UNSUPPORTED_FILE',400)
    except RescueError:raise
    except Exception as ex:raise RescueError('PPT/Excel读取失败，文件可能受保护或已损坏，请另存后再试。','STRUCTURED_FILE_ERROR',422) from ex
    # Image-heavy PPT decks can yield a valid outline but almost no text; surface that limitation so the Agent does not overclaim.
    thin = ext in ('.pptx','.pptm') and (len(text.strip()) / max(1, len(outline)) < 60)
    note = ('提取到的幻灯片文字偏少，可能有较多内容位于图片/图表中。当前只能做文字层初筛，请补充关键页截图或导出 PDF 后进行视觉复核。' if thin else '')
    return {'title':Path(name).stem,'text':text[:80000],'outline':outline,
            'chars':len(text),'source':name,'uploaded':True,'truncated':len(text)>80000,
            'thin':thin,'note':note}

def get_transcribe_state():
    from local_speech_bridge import state
    local=state()
    with _TRANS_LOCK:
        return {'configured':bool(_TRANSCRIBE['key']) or local['available'],
                'apiConfigured':bool(_TRANSCRIBE['key']) and _TRANSCRIBE['enabled'],'local':local,
                'enabled':_TRANSCRIBE['enabled'] or local['available'],
                'base':_TRANSCRIBE['base'],'model':_TRANSCRIBE['model'],
                'note':local['note']+' 独立 API Key 仅放本机进程内存。'}

def set_transcribe_config(data):
    base=str(data.get('base') or 'https://api.openai.com/v1').rstrip('/')
    parsed=urlparse(base)
    if parsed.scheme!='https' or (parsed.hostname not in ('api.openai.com',)):
        raise RescueError('目前只支持 OpenAI 官方转写 API，地址必须是 https://api.openai.com/v1','AUDIO_PROVIDER_UNSUPPORTED',400)
    model=str(data.get('model') or 'gpt-4o-mini-transcribe')
    if model not in ('gpt-4o-mini-transcribe','gpt-4o-transcribe','whisper-1'):
        raise RescueError('请选择官方支持的语音转写模型','AUDIO_MODEL_UNSUPPORTED',400)
    key=str(data.get('key') or '').strip()
    if len(key)>1000 or '\n' in key:raise RescueError('密钥格式不正确','INVALID_KEY',400)
    enabled=bool(data.get('enabled',True))
    with _TRANS_LOCK:
        if key:_TRANSCRIBE['key']=key
        _TRANSCRIBE.update({'base':base,'model':model,'enabled':enabled})
    return get_transcribe_state()

def transcribe_video_file(file,explicit):
    if not explicit:return {'text':'','status':'not-requested'}
    with _TRANS_LOCK:cfg=dict(_TRANSCRIBE)
    if not cfg['enabled'] or not cfg['key']:
        from local_speech_bridge import transcribe
        return transcribe(file)
    with tempfile.TemporaryDirectory(prefix='sy-video-audio-') as tmp:
        wav=Path(tmp)/'audio.mp3'
        command=['ffmpeg','-nostdin','-y','-v','error','-i',str(file),'-vn','-ac','1','-ar','16000',
                 '-codec:a','libmp3lame','-b:a','48k',str(wav)]
        try:
            p=subprocess.run(command,capture_output=True,timeout=60)
            if p.returncode!=0 or not wav.is_file() or wav.stat().st_size<256:
                return {'text':'','status':'no-audio-track','note':'文件没有可识别的音轨。'}
            if wav.stat().st_size>24*1024*1024:
                return {'text':'','status':'audio-too-large','note':'音频超过24MB，请先裁剪视频。'}
            with wav.open('rb') as audio:
                r=requests.post(cfg['base']+'/audio/transcriptions',
                    headers={'Authorization':'Bearer '+cfg['key']},
                    files={'file':('clip.mp3',audio,'audio/mpeg')},
                    data={'model':cfg['model']},timeout=100)
            if not r.ok:return {'text':'','status':'provider-error','note':'转写服务请求失败 HTTP '+str(r.status_code)+'。请检查模型、权限、余额。'}
            text=str(r.json().get('text') or '').strip()
            return {'text':text[:32000],'status':'done' if text else 'empty','model':cfg['model']}
        except (ValueError,requests.RequestException,OSError,subprocess.TimeoutExpired) as e:
            return {'text':'','status':'error','note':'转写未完成，请核对网络和服务商。'}
